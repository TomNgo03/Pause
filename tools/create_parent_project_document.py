from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK


OUTPUT = "/Users/tungngo/Documents/ChatGPT/Huy_vinschool/Tai_lieu_gioi_thieu_du_an_Pause_danh_cho_phu_huynh.docx"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.add_run(f"{index}.  ").bold = True
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, "332768")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Inches(widths[i])
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_fill(cell, "F5F3FB" if r_index % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.add_run(value)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(11.2)
normal.font.color.rgb = RGBColor(30, 30, 34)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(29)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
title_style_border = title_style_ppr.find(qn("w:pBdr"))
if title_style_border is not None:
    title_style_ppr.remove(title_style_border)

for level, size in ((1, 18), (2, 14), (3, 12)):
    style = styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    style.paragraph_format.space_after = Pt(6)

body = styles["Body Text"]
body.font.name = "Aptos"
body.font.size = Pt(11.2)
body.font.color.rgb = RGBColor(30, 30, 34)
body.paragraph_format.space_after = Pt(8)
body.paragraph_format.line_spacing = 1.12

for style_name in ("List Bullet", "List Number"):
    style = styles[style_name]
    style.font.name = "Aptos"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(30, 30, 34)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.08

if "Small Note" not in [s.name for s in styles]:
    small = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
else:
    small = styles["Small Note"]
small.font.name = "Aptos"
small.font.size = Pt(9.5)
small.font.color.rgb = RGBColor(92, 92, 102)
small.paragraph_format.space_after = Pt(5)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(68)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PAUSE")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(17)
r.font.color.rgb = RGBColor(84, 62, 160)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Tài liệu giới thiệu dự án Pause dành cho phụ huynh")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_before = Pt(10)
subtitle.paragraph_format.space_after = Pt(26)
run = subtitle.add_run("Ứng dụng hỗ trợ học sinh sử dụng thiết bị số có chủ đích và xây dựng khả năng tự quản lý")
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(70, 70, 78)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.paragraph_format.left_indent = Inches(0.75)
intro.paragraph_format.right_indent = Inches(0.75)
intro.paragraph_format.line_spacing = 1.2
intro.add_run(
    "Tài liệu này giúp phụ huynh hiểu Pause giải quyết vấn đề gì, học sinh trải nghiệm ứng dụng ra sao, "
    "dự án bảo vệ quyền riêng tư như thế nào và thành công sẽ được đánh giá bằng những dấu hiệu cụ thể nào."
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(65)
r = p.add_run("Bản trình bày định hướng dự án")
r.bold = True
r.font.size = Pt(11)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.add_run("Tháng 9 năm 2026").font.color.rgb = RGBColor(90, 90, 98)

page_break(doc)

add_heading(doc, "Tóm tắt dành cho phụ huynh", 1)
add_body(doc, "Pause là một ứng dụng di động được xây dựng cho học sinh đang gặp khó khăn trong việc duy trì sự tập trung trước mạng xã hội, trò chơi, video ngắn và các nội dung giải trí liên tục. Mục tiêu của dự án không phải là cấm học sinh dùng điện thoại, mà là giúp các em nhận biết mục đích sử dụng, tự chọn một khoảng thời gian hợp lý và từng bước hình thành khả năng kiểm soát bản thân.")
add_body(doc, "Ý tưởng trung tâm của Pause là tạo ra một khoảng dừng có ý nghĩa trước hành vi sử dụng thiết bị theo quán tính. Khi chuẩn bị bắt đầu một hoạt động, học sinh được khuyến khích xác định mình đang muốn làm gì, cần bao nhiêu thời gian và kết quả mong muốn là gì. Trong phiên tập trung, tiến độ được thể hiện bằng một hành trình vũ trụ nhẹ nhàng: phi thuyền di chuyển quanh hành tinh theo thời gian đã hoàn thành. Sau phiên, học sinh tự phản hồi về mức độ tập trung và điều cần điều chỉnh.")
add_body(doc, "Dự án coi sự tiến bộ là khả năng ra quyết định tốt hơn, chứ không đơn thuần là giảm số phút dùng điện thoại. Vì vậy, Pause ưu tiên mục tiêu thực tế, phản tư cá nhân, động viên tích cực, quyền riêng tư và sự tự nguyện. Ứng dụng không được mô tả như một công cụ điều trị nghiện, không thay thế tư vấn tâm lý và không dùng cảm giác tội lỗi để ép buộc hành vi.")

add_heading(doc, "Điều phụ huynh cần biết ngay", 2)
add_bullets(doc, [
    "Pause hướng tới việc rèn luyện thói quen tự chủ, không phải giám sát bí mật hoặc kiểm soát học sinh từ xa.",
    "Ứng dụng phù hợp nhất khi được sử dụng như một công cụ hỗ trợ học tập và trò chuyện trong gia đình.",
    "Các yếu tố trò chơi như hành tinh, huy hiệu và cột mốc chỉ nhằm ghi nhận nỗ lực; chúng không được thiết kế để tạo lệ thuộc.",
    "Thông tin cá nhân và hoạt động của bạn bè phải được chia sẻ theo lựa chọn rõ ràng của người dùng.",
    "Giá trị quan trọng nhất của dự án là quá trình học sinh nghiên cứu vấn đề, thử nghiệm giải pháp, lắng nghe phản hồi và cải tiến sản phẩm."
])

add_heading(doc, "Mục lục nội dung", 2)
add_numbered(doc, [
    "Bối cảnh và vấn đề dự án quan tâm",
    "Tầm nhìn và mục tiêu của Pause",
    "Hành trình trải nghiệm của một học sinh",
    "Các nhóm tính năng chính",
    "Giá trị giáo dục và tác động kỳ vọng",
    "Nguyên tắc đạo đức, an toàn và quyền riêng tư",
    "Cách thử nghiệm và đánh giá hiệu quả",
    "Vai trò của phụ huynh và nhà trường",
    "Phạm vi một tháng và định hướng tiếp theo",
    "Câu hỏi thường gặp"
])

page_break(doc)

add_heading(doc, "1 Bối cảnh và vấn đề dự án quan tâm", 1)
add_heading(doc, "Thiết bị số vừa hữu ích vừa dễ tạo thói quen tự động", 2)
add_body(doc, "Điện thoại và máy tính giúp học sinh tiếp cận kiến thức, trao đổi với giáo viên, làm bài tập và duy trì các mối quan hệ. Tuy nhiên, cùng một thiết bị cũng chứa nhiều nội dung được thiết kế để giữ sự chú ý càng lâu càng tốt. Video ngắn, thông báo liên tục, phần thưởng trong trò chơi và dòng nội dung không có điểm kết thúc khiến việc dừng lại trở nên khó khăn, đặc biệt khi học sinh đang mệt, căng thẳng hoặc chưa có kế hoạch rõ ràng.")
add_body(doc, "Một biểu hiện phổ biến là học sinh mở điện thoại chỉ để kiểm tra một thông báo nhưng tiếp tục sử dụng lâu hơn dự định. Vấn đề không nhất thiết bắt đầu từ ý định xấu hay thiếu kỷ luật. Nhiều hành vi xảy ra theo quán tính, trước khi người dùng kịp nhận ra mình đang lựa chọn điều gì. Hệ quả có thể là trì hoãn công việc, ngủ muộn, phân tán trong giờ học, khó hoàn thành nhiệm vụ dài và cảm giác hối tiếc sau khi sử dụng thiết bị.")

add_heading(doc, "Giới hạn cứng chưa giải quyết toàn bộ nguyên nhân", 2)
add_body(doc, "Các công cụ khóa ứng dụng hoặc đặt giới hạn thời gian có thể hữu ích trong một số trường hợp. Tuy nhiên, nếu học sinh chỉ tuân thủ vì bị chặn, các em có thể chưa học được cách nhận diện cám dỗ, cân nhắc ưu tiên và tự đưa ra quyết định. Khi giới hạn được gỡ bỏ, hành vi cũ có thể quay lại. Một số học sinh còn tìm cách vượt qua hệ thống hoặc cảm thấy bị kiểm soát, từ đó làm giảm sự hợp tác.")
add_body(doc, "Pause xuất phát từ câu hỏi khác: làm thế nào để hỗ trợ học sinh dừng lại trong vài giây, gọi tên ý định của mình và chủ động lựa chọn? Đây là một mục tiêu khiêm tốn nhưng có ý nghĩa lâu dài, bởi khả năng tự quản lý không chỉ cần cho việc dùng điện thoại mà còn cần cho học tập, nghỉ ngơi và cuộc sống trưởng thành.")

add_heading(doc, "Phạm vi vấn đề mà dự án lựa chọn", 2)
add_table(doc, ["Trong phạm vi", "Ngoài phạm vi"], [
    ("Hỗ trợ nhận biết thói quen sử dụng thiết bị", "Chẩn đoán hoặc điều trị nghiện công nghệ"),
    ("Giúp lập phiên tập trung có mục tiêu", "Bảo đảm điểm số hoặc thành tích học tập"),
    ("Khuyến khích tự phản hồi và điều chỉnh", "Thay thế phụ huynh, giáo viên hoặc chuyên gia"),
    ("Ghi nhận tiến bộ theo cách tích cực", "Theo dõi bí mật hoặc trừng phạt người dùng")
], widths=[3.35, 3.35])

page_break(doc)

add_heading(doc, "2 Tầm nhìn và mục tiêu của Pause", 1)
add_heading(doc, "Tầm nhìn", 2)
add_body(doc, "Pause hướng đến một tương lai trong đó học sinh có thể sử dụng công nghệ như một công cụ phục vụ mục tiêu của mình, thay vì bị cuốn theo thiết bị một cách vô thức. Ứng dụng không cố gắng loại bỏ giải trí. Giải trí là một phần bình thường của cuộc sống; điều quan trọng là người dùng biết mình đang lựa chọn, lựa chọn trong bao lâu và lựa chọn đó có phù hợp với hoàn cảnh hiện tại hay không.")

add_heading(doc, "Ba mục tiêu thay đổi hành vi", 2)
add_numbered(doc, [
    "Tăng nhận thức: giúp học sinh nhận ra thời điểm mình sắp hành động theo quán tính và xác định mục đích trước khi bắt đầu.",
    "Tăng khả năng tự quản lý: hỗ trợ đặt thời gian phù hợp, hoàn thành phiên đã cam kết và điều chỉnh mục tiêu dựa trên trải nghiệm thực tế.",
    "Tạo động lực lành mạnh: ghi nhận sự nhất quán và nỗ lực mà không biến thành cuộc đua gây áp lực hoặc cảm giác thất bại."
])

add_heading(doc, "Các nguyên tắc thiết kế trải nghiệm", 2)
add_table(doc, ["Nguyên tắc", "Ý nghĩa đối với học sinh"], [
    ("Có chủ đích", "Mỗi phiên bắt đầu bằng một mục tiêu rõ ràng, không chỉ bằng việc bấm đồng hồ."),
    ("Linh hoạt", "Học sinh chọn thời lượng phù hợp thay vì bị áp dụng một giới hạn giống nhau."),
    ("Tích cực", "Ứng dụng ghi nhận nỗ lực và tiến bộ; không dùng ngôn ngữ làm người dùng xấu hổ."),
    ("Dễ hiểu", "Màn hình chính và phiên tập trung phải đơn giản, ít lựa chọn gây phân tâm."),
    ("Tự nguyện", "Người dùng có quyền quyết định dữ liệu nào được lưu hoặc chia sẻ."),
    ("Dần độc lập", "Thành công lâu dài là khi học sinh có thể tự quản lý tốt hơn, kể cả khi không mở Pause.")
], widths=[1.55, 5.15])

add_heading(doc, "Đối tượng sử dụng", 2)
add_body(doc, "Đối tượng chính là học sinh trung học muốn cải thiện khả năng tập trung, quản lý thời gian và hiểu rõ hơn thói quen số của mình. Ứng dụng cũng có thể hỗ trợ sinh viên hoặc người trẻ, nhưng nội dung và cách thử nghiệm ban đầu nên tập trung vào nhóm học sinh để dự án có phạm vi rõ ràng và phản hồi đủ sâu.")

page_break(doc)

add_heading(doc, "3 Hành trình trải nghiệm của một học sinh", 1)
add_body(doc, "Để minh họa, có thể hình dung một học sinh tên Minh. Sau giờ học, Minh cần làm bài Toán trong 30 phút nhưng thường mở video ngắn khi thấy bài khó. Pause không tự động phán xét hoặc khóa toàn bộ điện thoại. Ứng dụng hướng Minh đi qua một chuỗi quyết định ngắn và dễ hiểu.")

add_heading(doc, "Bước 1 Tạo tài khoản cá nhân", 2)
add_body(doc, "Khi mở ứng dụng lần đầu, Minh tạo tài khoản hoặc đăng nhập. Tài khoản giúp lưu tiến độ và hồ sơ thành tích trên nhiều lần sử dụng. Quy trình cần giải thích rõ thông tin nào được thu thập và không yêu cầu dữ liệu không cần thiết.")

add_heading(doc, "Bước 2 Xác định ý định", 2)
add_body(doc, "Trước phiên tập trung, Minh chọn hoặc nhập mục tiêu, chẳng hạn hoàn thành năm bài Toán. Việc gọi tên một kết quả cụ thể giúp chuyển từ suy nghĩ chung chung như phải học sang một hành động có thể bắt đầu ngay.")

add_heading(doc, "Bước 3 Chọn thời lượng thực tế", 2)
add_body(doc, "Minh chọn 25 hoặc 30 phút dựa trên thời gian và năng lượng hiện có. Pause khuyến khích một cam kết vừa sức. Một phiên ngắn được hoàn thành nghiêm túc có giá trị hơn một kế hoạch dài nhưng liên tục bị bỏ dở.")

add_heading(doc, "Bước 4 Bắt đầu hành trình tập trung", 2)
add_body(doc, "Trong phiên, màn hình trở nên yên tĩnh. Một hành tinh nằm ở trung tâm, thời gian còn lại được hiển thị rõ và một phi thuyền di chuyển quanh quỹ đạo theo phần trăm thời gian đã hoàn thành. Hình ảnh này biến tiến độ trừu tượng thành một hành trình dễ cảm nhận, nhưng không thêm quá nhiều nút bấm hoặc hiệu ứng khiến học sinh tiếp tục nhìn vào điện thoại.")

add_heading(doc, "Bước 5 Hoàn thành và phản tư", 2)
add_body(doc, "Khi phiên kết thúc, Minh ghi nhận mình đã hoàn thành nhiệm vụ ở mức nào, mức độ tập trung ra sao và điều gì gây phân tâm. Đây không phải một bài kiểm tra. Mục tiêu là giúp Minh nhận ra mẫu hành vi, ví dụ học tốt hơn vào đầu buổi tối hoặc cần chia nhiệm vụ lớn thành các phần nhỏ.")

add_heading(doc, "Bước 6 Ghi nhận cột mốc", 2)
add_body(doc, "Sau nhiều phiên, Minh có thể nhận huy hiệu hoặc khám phá một cột mốc mới trong hành trình vũ trụ. Thành tích có thể xuất hiện trên hồ sơ cá nhân và được chia sẻ với bạn bè nếu Minh chủ động cho phép. Trọng tâm là sự nhất quán, không phải cạnh tranh về số giờ.")

page_break(doc)

add_heading(doc, "4 Các nhóm tính năng chính", 1)
add_heading(doc, "Tài khoản và hồ sơ cá nhân", 2)
add_body(doc, "Mỗi người dùng có một hồ sơ riêng để lưu mục tiêu, phiên tập trung, cột mốc và huy hiệu. Hồ sơ giúp học sinh nhìn thấy hành trình của chính mình theo thời gian. Các thao tác cơ bản gồm tạo tài khoản, đăng nhập, đăng xuất và khôi phục mật khẩu.")

add_heading(doc, "Lập kế hoạch phiên tập trung", 2)
add_body(doc, "Học sinh chọn mục đích và thời lượng trước khi bắt đầu. Ứng dụng có thể gợi ý mục tiêu nhỏ hơn khi nhiệm vụ quá rộng. Ví dụ, thay vì học Hóa, một kế hoạch rõ ràng hơn là đọc mục 2, ghi năm ý chính và làm ba câu hỏi trong 25 phút.")

add_heading(doc, "Trợ lý lập kế hoạch", 2)
add_body(doc, "Một trợ lý trong ứng dụng có thể hỗ trợ học sinh biến mục tiêu thành kế hoạch phù hợp với thời gian và trạng thái hiện tại. Trợ lý chỉ đưa ra gợi ý; học sinh vẫn là người quyết định. Nội dung cần tránh chẩn đoán tâm lý, tránh khẳng định chắc chắn và luôn cho phép người dùng sửa kế hoạch.")

add_heading(doc, "Phiên tập trung theo hành trình vũ trụ", 2)
add_body(doc, "Giao diện phiên tập trung sử dụng hình ảnh hành tinh, quỹ đạo và phi thuyền để thể hiện tiến độ. Khi thời gian trôi qua, phi thuyền tiến dần quanh vòng tròn. Màn hình giữ thiết kế tối giản để hỗ trợ tập trung thay vì trở thành một nội dung giải trí mới.")

add_heading(doc, "Theo dõi tiến bộ và phản tư", 2)
add_body(doc, "Sau mỗi phiên, học sinh có thể đánh giá mức độ tập trung và kết quả. Theo thời gian, ứng dụng tổng hợp số phiên hoàn thành, thời gian tập trung và tính nhất quán. Các số liệu này nên được diễn giải cẩn thận: nhiều giờ hơn không mặc nhiên đồng nghĩa với học tốt hơn.")

add_heading(doc, "Huy hiệu, cột mốc và bộ sưu tập", 2)
add_body(doc, "Huy hiệu có thể ghi nhận phiên đầu tiên, một tuần duy trì đều đặn, nhiều lần hoàn thành mục tiêu hoặc sự trở lại sau khi gián đoạn. Việc ghi nhận khả năng quay lại là quan trọng, vì thói quen tốt không cần một chuỗi hoàn hảo mới có giá trị.")

add_heading(doc, "Bạn bè và thử thách chung", 2)
add_body(doc, "Học sinh có thể kết nối với bạn bè, xem những thành tích mà bạn chọn công khai, gửi lời động viên hoặc tham gia thử thách chung. Thiết kế cần tránh bảng xếp hạng tuyệt đối theo số giờ, bởi điều kiện học tập và nhu cầu nghỉ ngơi của mỗi người khác nhau.")

page_break(doc)

add_heading(doc, "5 Giá trị giáo dục và tác động kỳ vọng", 1)
add_heading(doc, "Khả năng tự nhận thức", 2)
add_body(doc, "Thông qua việc xác định ý định trước phiên và phản hồi sau phiên, học sinh có cơ hội quan sát cách mình làm việc. Các em dần nhận ra lúc nào dễ mất tập trung, thời lượng nào phù hợp và môi trường nào giúp hoàn thành nhiệm vụ tốt hơn.")

add_heading(doc, "Kỹ năng lập kế hoạch", 2)
add_body(doc, "Pause khuyến khích chuyển một nhiệm vụ lớn thành mục tiêu có phạm vi cụ thể. Kỹ năng này hữu ích trong học tập, hoạt động ngoại khóa và các dự án dài hạn. Học sinh học cách ước lượng thời gian, bắt đầu từ một bước nhỏ và điều chỉnh sau trải nghiệm thực tế.")

add_heading(doc, "Khả năng trì hoãn phần thưởng", 2)
add_body(doc, "Một phiên tập trung là cam kết dành sự chú ý cho nhiệm vụ trong một khoảng thời gian đã chọn. Việc hoàn thành nhiều cam kết nhỏ có thể giúp học sinh luyện khả năng không phản ứng ngay với mọi thông báo hoặc cám dỗ. Tuy nhiên, dự án không nên tuyên bố rằng một ứng dụng đơn lẻ có thể thay đổi hoàn toàn hành vi.")

add_heading(doc, "Mối quan hệ lành mạnh hơn với công nghệ", 2)
add_body(doc, "Mục tiêu không phải là càng ít thời gian màn hình càng tốt. Một giờ học trực tuyến và một giờ lướt nội dung không chủ đích có ý nghĩa khác nhau. Pause hướng tới chất lượng quyết định: người dùng biết mình đang làm gì, vì sao làm và khi nào nên dừng.")

add_heading(doc, "Tác động kỳ vọng và dấu hiệu quan sát", 2)
add_table(doc, ["Kết quả kỳ vọng", "Dấu hiệu có thể quan sát"], [
    ("Ít sử dụng theo quán tính", "Học sinh chủ động xác định mục đích trước khi dùng thiết bị."),
    ("Kế hoạch thực tế hơn", "Mục tiêu phiên cụ thể và tỷ lệ hoàn thành tăng dần."),
    ("Hiểu bản thân hơn", "Học sinh mô tả được nguyên nhân gây phân tâm và cách điều chỉnh."),
    ("Duy trì tốt hơn", "Học sinh quay lại sau một phiên không thành công thay vì bỏ cuộc."),
    ("Giao tiếp tích cực", "Bạn bè động viên nhau mà không chế giễu hoặc gây áp lực thành tích.")
], widths=[2.25, 4.45])

add_body(doc, "Điểm số học tập có thể được xem như một thông tin tham khảo nếu học sinh và gia đình đồng ý, nhưng không nên được dùng làm bằng chứng duy nhất. Điểm số chịu ảnh hưởng của nhiều yếu tố ngoài ứng dụng, và thử nghiệm nhỏ không đủ để khẳng định quan hệ nguyên nhân kết quả.")

page_break(doc)

add_heading(doc, "6 Nguyên tắc đạo đức, an toàn và quyền riêng tư", 1)
add_heading(doc, "Không dùng sự xấu hổ để thay đổi hành vi", 2)
add_body(doc, "Pause không nên gọi người dùng là lười biếng, thiếu kỷ luật hoặc thất bại. Nếu một phiên bị dừng, ứng dụng có thể hỏi điều gì đã xảy ra và đề nghị bắt đầu lại với mục tiêu nhỏ hơn. Ngôn ngữ cần bình tĩnh, tôn trọng và tập trung vào hành động tiếp theo.")

add_heading(doc, "Không tạo một dạng lệ thuộc mới", 2)
add_body(doc, "Một ứng dụng giảm phân tâm không nên giữ học sinh trên màn hình bằng thông báo dồn dập, phần thưởng ngẫu nhiên hoặc chuỗi ngày gây sợ mất thành tích. Hiệu ứng hình ảnh phải vừa đủ để ghi nhận tiến bộ. Sau khi bắt đầu phiên, người dùng không cần liên tục tương tác với Pause.")

add_heading(doc, "Thu thập dữ liệu tối thiểu", 2)
add_body(doc, "Chỉ nên lưu những thông tin cần thiết để tài khoản và tính năng hoạt động. Học sinh và phụ huynh cần được biết dữ liệu nào được lưu, mục đích sử dụng, ai có thể xem và cách yêu cầu xóa. Không nên bán dữ liệu cá nhân hoặc dùng dữ liệu học sinh cho quảng cáo hành vi.")

add_heading(doc, "Chia sẻ theo lựa chọn", 2)
add_body(doc, "Thành tích và hoạt động bạn bè phải mặc định ở mức riêng tư hợp lý. Người dùng cần chủ động chọn nội dung công khai. Các thông tin nhạy cảm như nội dung phản tư, điểm số, thời gian biểu chi tiết hoặc cảm xúc cá nhân không nên tự động xuất hiện trên hồ sơ bạn bè.")

add_heading(doc, "Bảo vệ người dùng chưa thành niên", 2)
add_body(doc, "Vì đối tượng chính là học sinh, dự án cần đặc biệt thận trọng với thông tin nhận dạng, kết nối bạn bè và nội dung do người dùng tạo. Giai đoạn thử nghiệm nên giới hạn trong nhóm nhỏ có sự đồng thuận phù hợp, có quy tắc ứng xử rõ ràng và có người phụ trách tiếp nhận báo cáo.")

add_heading(doc, "Giới hạn của trợ lý lập kế hoạch", 2)
add_body(doc, "Trợ lý chỉ hỗ trợ tổ chức nhiệm vụ và thời gian. Nếu người dùng đề cập đến tự gây hại, khủng hoảng tâm lý hoặc vấn đề sức khỏe nghiêm trọng, ứng dụng không nên tự xử lý như chuyên gia. Thay vào đó, cần khuyến khích liên hệ người lớn đáng tin cậy hoặc dịch vụ hỗ trợ phù hợp.")

add_heading(doc, "Cam kết đạo đức đề xuất", 2)
add_bullets(doc, [
    "Không theo dõi bí mật và không chia sẻ dữ liệu nếu chưa có lựa chọn rõ ràng.",
    "Không coi số giờ tập trung là thước đo giá trị của học sinh.",
    "Không sử dụng bảng xếp hạng gây áp lực hoặc làm lộ thông tin riêng tư.",
    "Không tuyên bố Pause điều trị nghiện, trầm cảm, lo âu hoặc rối loạn chú ý.",
    "Luôn cho phép người dùng chỉnh sửa mục tiêu, dừng phiên và xóa dữ liệu của mình.",
    "Đánh giá tác động bằng phản hồi trung thực, kể cả khi kết quả không như kỳ vọng."
])

page_break(doc)

add_heading(doc, "7 Cách thử nghiệm và đánh giá hiệu quả", 1)
add_heading(doc, "Mục đích của thử nghiệm ban đầu", 2)
add_body(doc, "Trong giai đoạn đầu, mục tiêu không phải chứng minh ứng dụng hiệu quả với mọi học sinh. Mục tiêu là tìm hiểu liệu trải nghiệm có dễ hiểu, có giúp người dùng dừng lại và suy nghĩ, có hỗ trợ hoàn thành nhiệm vụ và có tạo ra tác dụng không mong muốn hay không.")

add_heading(doc, "Nhóm thử nghiệm đề xuất", 2)
add_body(doc, "Nên mời khoảng 10 đến 15 học sinh dùng phiên bản thử nghiệm trong thời gian ngắn. Việc tham gia cần tự nguyện, được giải thích rõ mục đích và có phương án xin sự đồng thuận phù hợp. Người tham gia được quyền rút lui mà không chịu bất kỳ bất lợi nào.")

add_heading(doc, "Những thông tin nên thu thập", 2)
add_table(doc, ["Nhóm thông tin", "Ví dụ", "Mục đích"], [
    ("Khả năng sử dụng", "Có hiểu cách bắt đầu phiên không", "Phát hiện điểm gây nhầm lẫn"),
    ("Hành vi phiên", "Số phiên bắt đầu và hoàn thành", "Quan sát mức độ duy trì"),
    ("Tự đánh giá", "Mức tập trung trước và sau", "Hiểu trải nghiệm chủ quan"),
    ("Phản hồi mở", "Điều hữu ích hoặc gây khó chịu", "Tìm tác động ngoài dự kiến"),
    ("Thói quen số", "Thời gian dùng ứng dụng gây phân tâm", "Quan sát xu hướng, không phán xét")
], widths=[1.45, 2.65, 2.6])

add_heading(doc, "Cách diễn giải kết quả", 2)
add_body(doc, "Dữ liệu của nhóm nhỏ chỉ giúp định hướng cải tiến. Nếu thời gian màn hình giảm, chưa thể kết luận Pause là nguyên nhân duy nhất. Nếu không giảm, ứng dụng vẫn có thể tạo giá trị nếu học sinh sử dụng thiết bị có mục đích hơn. Báo cáo nên trình bày cả kết quả tích cực, phản hồi tiêu cực, hạn chế và những điều chưa biết.")

add_heading(doc, "Tiêu chí thành công ban đầu", 2)
add_bullets(doc, [
    "Phần lớn người thử nghiệm có thể tự tạo và hoàn thành một phiên mà không cần hướng dẫn trực tiếp.",
    "Người dùng hiểu rằng Pause hỗ trợ tự quản lý chứ không phải một công cụ trừng phạt.",
    "Câu hỏi phản tư được đánh giá là ngắn gọn, phù hợp và không gây phán xét.",
    "Không xuất hiện vấn đề nghiêm trọng về quyền riêng tư hoặc tương tác bạn bè.",
    "Nhóm phát triển xác định được các cải tiến cụ thể dựa trên phản hồi thực tế."
])

page_break(doc)

add_heading(doc, "8 Vai trò của phụ huynh và nhà trường", 1)
add_heading(doc, "Phụ huynh là người đồng hành", 2)
add_body(doc, "Pause phát huy giá trị tốt nhất khi phụ huynh coi ứng dụng là cơ hội trò chuyện chứ không phải công cụ giám sát. Có thể hỏi con điều gì giúp con tập trung, thời lượng nào phù hợp hoặc con muốn gia đình hỗ trợ bằng cách nào. Những câu hỏi mở thường tạo hợp tác tốt hơn việc chỉ yêu cầu giảm số giờ sử dụng điện thoại.")

add_heading(doc, "Những cách hỗ trợ phù hợp", 2)
add_bullets(doc, [
    "Cùng thống nhất thời điểm học, nghỉ và giải trí thay vì áp đặt liên tục.",
    "Khuyến khích mục tiêu nhỏ, cụ thể và có thể hoàn thành.",
    "Ghi nhận nỗ lực và khả năng quay lại sau khi mất tập trung.",
    "Tôn trọng không gian phản tư riêng tư của học sinh.",
    "Làm gương bằng cách giảm thông báo hoặc để điện thoại sang một bên trong thời gian gia đình.",
    "Tìm hỗ trợ chuyên môn nếu việc sử dụng thiết bị liên quan đến mất ngủ kéo dài, khủng hoảng tâm lý hoặc suy giảm chức năng nghiêm trọng."
])

add_heading(doc, "Những điều nên tránh", 2)
add_bullets(doc, [
    "Dùng số liệu trong ứng dụng để so sánh con với bạn bè hoặc anh chị em.",
    "Yêu cầu con công khai toàn bộ lịch sử hoạt động và nội dung phản tư.",
    "Coi một phiên không hoàn thành là bằng chứng về thái độ hoặc phẩm chất.",
    "Biến huy hiệu và chuỗi ngày thành điều kiện thưởng phạt quá lớn.",
    "Kỳ vọng ứng dụng tự giải quyết mọi vấn đề học tập hoặc sức khỏe tinh thần."
])

add_heading(doc, "Vai trò của nhà trường", 2)
add_body(doc, "Nhà trường có thể hỗ trợ bằng cách cung cấp bối cảnh thử nghiệm an toàn, góp ý về nội dung giáo dục và giúp học sinh hiểu về sức khỏe số. Nếu dự án được thử nghiệm trong trường, cần có người phụ trách rõ ràng, phạm vi dữ liệu giới hạn và cơ chế tiếp nhận phản hồi. Giáo viên không nên dùng dữ liệu Pause như một hình thức chấm điểm hoặc kỷ luật.")

page_break(doc)

add_heading(doc, "9 Phạm vi một tháng và định hướng tiếp theo", 1)
add_body(doc, "Vì dự án có thời gian phát triển ngắn, phiên bản đầu tiên nên ưu tiên một hành trình hoàn chỉnh và đáng tin cậy thay vì cố gắng hoàn thiện mọi ý tưởng. Một sản phẩm nhỏ nhưng có thể sử dụng, được thử nghiệm và có bài học rõ ràng sẽ thuyết phục hơn một danh sách tính năng lớn chưa được kiểm chứng.")

add_heading(doc, "Ưu tiên trong tháng đầu", 2)
add_table(doc, ["Giai đoạn", "Trọng tâm", "Kết quả mong đợi"], [
    ("Tuần 1", "Làm rõ nhu cầu và luồng sử dụng", "Học sinh hiểu cách tạo tài khoản và bắt đầu phiên"),
    ("Tuần 2", "Hoàn thiện phiên tập trung", "Đồng hồ, ý định và hành trình phi thuyền hoạt động ổn định"),
    ("Tuần 3", "Tiến bộ và phản tư", "Lưu kết quả, hiển thị cột mốc và nhận phản hồi sau phiên"),
    ("Tuần 4", "Thử nghiệm và cải tiến", "Có phản hồi người dùng, sửa lỗi chính và chuẩn bị phần trình bày")
], widths=[1.05, 2.75, 2.9])

add_heading(doc, "Tính năng nên được xem là định hướng tiếp theo", 2)
add_body(doc, "Các chức năng xã hội sâu hơn, nhiều hành tinh, thử thách nhóm, cá nhân hóa nâng cao và trợ lý lập kế hoạch phong phú có thể được phát triển sau khi trải nghiệm cốt lõi được kiểm chứng. Việc phân biệt rõ phiên bản hiện tại và định hướng tương lai giúp dự án trung thực, có trọng tâm và dễ đánh giá.")

add_heading(doc, "Giá trị của dự án đối với quá trình trưởng thành của học sinh", 2)
add_body(doc, "Dự án cho phép học sinh kết hợp trải nghiệm cá nhân với quan sát một vấn đề xã hội, sau đó chuyển ý tưởng thành sản phẩm có thể thử nghiệm. Trong quá trình đó, học sinh phải đặt câu hỏi, lựa chọn phạm vi, lắng nghe người dùng, cân nhắc đạo đức, chấp nhận giới hạn và sửa đổi thiết kế. Đây là những năng lực quan trọng đối với việc học khoa học máy tính và giải quyết vấn đề nói chung.")

add_heading(doc, "Cách trình bày dự án một cách trung thực", 2)
add_bullets(doc, [
    "Nêu rõ vấn đề bắt nguồn từ trải nghiệm và quan sát thực tế của học sinh.",
    "Giải thích vì sao chỉ khóa ứng dụng chưa đủ và Pause thử một hướng tiếp cận khác.",
    "Trình diễn một hành trình sử dụng hoàn chỉnh thay vì liệt kê mọi màn hình.",
    "Chia sẻ phản hồi thật của người thử nghiệm, bao gồm những điểm chưa tốt.",
    "Nói rõ đóng góp cá nhân, sự hỗ trợ đã nhận và bài học rút ra.",
    "Không phóng đại số người dùng, tác động xã hội hoặc khả năng của trợ lý."
])

page_break(doc)

add_heading(doc, "10 Câu hỏi thường gặp", 1)

faqs = [
    ("Pause có khóa mạng xã hội hoặc trò chơi không?", "Trọng tâm hiện tại của Pause là tạo ý thức và hỗ trợ phiên tập trung, không phải kiểm soát bằng khóa cứng. Một số giới hạn thiết bị có thể được cân nhắc sau, nhưng cần giữ quyền lựa chọn và tránh biến ứng dụng thành công cụ trừng phạt."),
    ("Ứng dụng có giúp tăng điểm số không?", "Pause có thể hỗ trợ thói quen lập kế hoạch và tập trung, nhưng không thể bảo đảm điểm số. Kết quả học tập còn phụ thuộc vào kiến thức nền, phương pháp học, giấc ngủ, sức khỏe, môi trường và nhiều yếu tố khác."),
    ("Vì sao lại sử dụng hành trình vũ trụ?", "Hành trình giúp học sinh cảm nhận tiến độ mà không cần đọc nhiều số liệu. Phi thuyền di chuyển quanh hành tinh theo thời gian hoàn thành, tạo cảm giác tiến triển nhẹ nhàng. Thiết kế vẫn phải đủ yên tĩnh để không gây phân tâm."),
    ("Huy hiệu có khiến học sinh lệ thuộc vào phần thưởng không?", "Rủi ro này có thật nếu phần thưởng được sử dụng quá mức. Vì vậy, Pause nên ghi nhận những hành vi có ý nghĩa như hoàn thành mục tiêu hoặc quay lại sau gián đoạn, đồng thời không dùng phần thưởng ngẫu nhiên hay cảnh báo gây sợ mất chuỗi."),
    ("Phụ huynh có xem được toàn bộ hoạt động của con không?", "Dự án ưu tiên quyền riêng tư và sự đồng thuận. Nếu có tính năng chia sẻ với phụ huynh trong tương lai, học sinh cần biết rõ nội dung nào được chia sẻ và có quyền kiểm soát phù hợp. Nội dung phản tư cá nhân không nên tự động công khai."),
    ("Thông tin của học sinh được sử dụng như thế nào?", "Chỉ nên lưu dữ liệu cần thiết cho tài khoản và tiến độ. Chính sách cần giải thích rõ mục đích, thời gian lưu, quyền truy cập và cách xóa dữ liệu. Dữ liệu học sinh không nên được bán hoặc dùng cho quảng cáo hành vi."),
    ("Trợ lý lập kế hoạch có quyết định thay học sinh không?", "Không. Trợ lý đề xuất một kế hoạch dựa trên mục tiêu và thời gian người dùng cung cấp. Học sinh có thể sửa hoặc bỏ đề xuất. Đây là công cụ hỗ trợ suy nghĩ, không phải người có thẩm quyền đưa ra quyết định."),
    ("Nếu học sinh liên tục hủy phiên thì sao?", "Ứng dụng nên phản hồi bằng câu hỏi nhẹ nhàng để tìm nguyên nhân và đề nghị mục tiêu nhỏ hơn. Việc hủy phiên không nên làm mất toàn bộ thành tích hoặc tạo cảm giác xấu hổ. Nếu khó khăn kéo dài và ảnh hưởng nghiêm trọng, gia đình nên tìm hiểu nguyên nhân rộng hơn."),
    ("Pause khác gì đồng hồ Pomodoro?", "Pomodoro chủ yếu chia thời gian thành các khoảng làm việc và nghỉ. Pause bổ sung bước xác định ý định, phản tư sau phiên, theo dõi cột mốc và nguyên tắc xây dựng khả năng tự quản lý trong bối cảnh sử dụng thiết bị số."),
    ("Dự án đã hoàn thiện chưa?", "Pause là một sản phẩm đang được phát triển và thử nghiệm. Trải nghiệm cốt lõi đã được xây dựng, nhưng các tính năng cần tiếp tục được kiểm tra với người dùng. Việc công khai giới hạn hiện tại là một phần quan trọng của cách làm dự án có trách nhiệm.")
]
for faq_index, (question, answer) in enumerate(faqs):
    if faq_index == 5:
        page_break(doc)
    add_heading(doc, question, 2)
    add_body(doc, answer)

page_break(doc)

add_heading(doc, "Kết luận", 1)
add_body(doc, "Pause đề xuất một cách tiếp cận nhân văn đối với vấn đề mất tập trung trong môi trường số. Thay vì chỉ dựa vào ngăn cấm, ứng dụng giúp học sinh tạo một khoảng dừng, xác định mục tiêu, thực hiện cam kết trong thời gian vừa sức và tự nhìn lại kết quả. Hành trình vũ trụ làm tiến độ trở nên trực quan, trong khi các nguyên tắc đạo đức giữ trải nghiệm tập trung vào sự trưởng thành của người dùng.")
add_body(doc, "Đối với phụ huynh, điều đáng quan tâm không chỉ là sản phẩm cuối cùng mà còn là cách học sinh xây dựng dự án: bắt đầu từ một vấn đề thật, xác định giới hạn, cân nhắc quyền riêng tư, thử nghiệm với người dùng và sẵn sàng thay đổi khi phản hồi cho thấy ý tưởng ban đầu chưa phù hợp. Đây là một quá trình học tập có chiều sâu và có thể tạo giá trị ngay cả khi sản phẩm vẫn còn đang hoàn thiện.")
add_body(doc, "Kỳ vọng hợp lý cho Pause là trở thành một công cụ nhỏ, dễ dùng và đáng tin cậy, giúp học sinh thực hiện những quyết định tốt hơn từng ngày. Nếu dự án giữ được sự trung thực, tôn trọng người dùng và tập trung vào bằng chứng từ thử nghiệm, Pause có thể vừa là một sản phẩm hữu ích vừa là một minh chứng rõ ràng cho tư duy giải quyết vấn đề có trách nhiệm.")

add_heading(doc, "Thông điệp trình bày ngắn", 2)
p = doc.add_paragraph(style="Body Text")
p.paragraph_format.left_indent = Inches(0.35)
p.paragraph_format.right_indent = Inches(0.35)
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(12)
r = p.add_run(
    "Pause không cố giành quyền kiểm soát chiếc điện thoại từ học sinh. Dự án giúp các em dừng lại đủ lâu để tự quyết định mình muốn dùng công nghệ như thế nào."
)
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(58, 42, 112)

note = doc.add_paragraph(style="Small Note")
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note.add_run("Tài liệu phục vụ giới thiệu định hướng dự án và trao đổi với phụ huynh. Các tính năng đang tiếp tục được thử nghiệm và cải tiến.")

# Header and footer
for sec in doc.sections:
    header = sec.header.paragraphs[0]
    header.text = "PAUSE   TÀI LIỆU DÀNH CHO PHỤ HUYNH"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor(110, 110, 118)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Pause   •   ")
    run.font.size = Pt(8.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

# Suppress header/footer on cover.
doc.sections[0].different_first_page_header_footer = True

doc.core_properties.title = "Tài liệu giới thiệu dự án Pause dành cho phụ huynh"
doc.core_properties.subject = "Giới thiệu mục tiêu, trải nghiệm, giá trị giáo dục và nguyên tắc đạo đức của dự án Pause"
doc.core_properties.author = "Nhóm dự án Pause"
doc.save(OUTPUT)
print(OUTPUT)
